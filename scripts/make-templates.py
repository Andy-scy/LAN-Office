# -*- coding: utf-8 -*-
"""生成空白 docx / xlsx / pptx 模板（优先用 python 库保证格式合法）。"""
import os, sys

OUT = os.path.join(os.path.dirname(__file__), "..", "server", "templates")
os.makedirs(OUT, exist_ok=True)


def make_docx():
    import docx
    d = docx.Document()
    d.add_paragraph("")
    d.save(os.path.join(OUT, "blank.docx"))


def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.save(os.path.join(OUT, "blank.xlsx"))


def make_pptx():
    import pptx
    p = pptx.Presentation()
    p.save(os.path.join(OUT, "blank.pptx"))


if __name__ == "__main__":
    results = []
    for name, fn in [("docx", make_docx), ("xlsx", make_xlsx), ("pptx", make_pptx)]:
        try:
            fn()
            results.append(f"{name}: OK")
        except Exception as e:
            results.append(f"{name}: FAIL ({e})")
    print("; ".join(results))
    sys.exit(0 if all(r.endswith("OK") for r in results) else 1)
